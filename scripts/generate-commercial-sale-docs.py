from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "Docs" / "Commercial Sales Templates"
LOGO = ROOT / "src" / "assets" / "logo.png"

BLACK = "202020"
GRAY = "666666"
LIGHT_GRAY = "E4E4E4"
VERY_LIGHT_GRAY = "F6F6F6"
PINK = "F672A2"
PALE_PINK = "FDEAF1"
WHITE = "FFFFFF"
VERSION = "Template version: September 2026"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LIGHT_GRAY, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=45, start=95, bottom=45, end=95) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_widths(table, widths: list[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = Inches(width)


def set_cell_text(cell, text: str, *, bold=False, color=BLACK, size=9.5, align=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def add_field_table(doc: Document, rows: list[tuple[str, str]], widths=(2.0, 4.8)):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for label, value in rows:
        row = table.add_row()
        prevent_row_split(row)
        set_cell_text(row.cells[0], label, bold=True, size=9.3)
        set_cell_shading(row.cells[0], VERY_LIGHT_GRAY)
        set_cell_text(row.cells[1], value, size=9.3)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
    set_table_widths(table, list(widths))
    return table


def add_grid_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    *,
    font_size=8.8,
    header_fill=PALE_PINK,
):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for idx, heading in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_text(cell, heading, bold=True, size=font_size)
        set_cell_shading(cell, header_fill)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            set_cell_text(cell, value, size=font_size)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_table_widths(table, widths)
    return table


def add_page_number(paragraph) -> None:
    paragraph.add_run("Page ")
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def configure_document(doc: Document, title: str) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.25)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    for style_name, size in (("Title", 17), ("Heading 1", 13), ("Heading 2", 11.5), ("Heading 3", 10.5)):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(9 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(4)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"Bloomjoy Sweets  |  bloomjoyusa.com  |  {VERSION}  |  ")
    r.font.name = "Arial"
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor.from_string(GRAY)
    add_page_number(p)
    for run in p.runs:
        run.font.name = "Arial"
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor.from_string(GRAY)

    doc.core_properties.title = title
    doc.core_properties.subject = "Bloomjoy reusable commercial machine agreement template"
    doc.core_properties.author = "TGPACI LLC dba Bloomjoy Sweets"
    doc.core_properties.keywords = "Bloomjoy, commercial machine, agreement, template"


def add_brand_block(doc: Document, descriptor: str) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Inches(0.85)
    table.columns[1].width = Inches(5.95)
    left, right = table.rows[0].cells
    left.width = Inches(0.85)
    right.width = Inches(5.95)
    p = left.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(LOGO), width=Inches(0.62))
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("BLOOMJOY SWEETS")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(PINK)
    p.add_run("\n")
    run = p.add_run(descriptor)
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    for cell in (left, right):
        set_cell_margins(cell, top=0, bottom=30, start=0, end=0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_title(doc: Document, title: str, subtitle: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor.from_string(BLACK)
    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(9)
        run = p.add_run(subtitle)
        run.italic = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor.from_string(GRAY)


def add_heading(doc: Document, text: str, level=1) -> None:
    doc.add_heading(text, level=level)


def add_paragraph(doc: Document, text: str, *, bold_lead: str | None = None, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.widow_control = True
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead) :])
    else:
        run = p.add_run(text)
        run.italic = italic
    return p


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def add_numbered_terms(doc: Document, terms: list[tuple[str, str]]) -> None:
    for heading, body in terms:
        p = doc.add_paragraph()
        p.paragraph_format.keep_together = False
        p.paragraph_format.widow_control = True
        run = p.add_run(f"{heading}  ")
        run.bold = True
        p.add_run(body)


def add_signature_table(doc: Document, left_label: str, right_label: str) -> None:
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    rows = [
        (left_label, right_label),
        ("By: __________________________________", "By: __________________________________"),
        ("Name: ________________________________", "Name: ________________________________"),
        ("Title: _________________________________", "Title: _________________________________"),
        ("Date: __________________________________", "Date: __________________________________"),
    ]
    for ridx, values in enumerate(rows):
        row = table.rows[ridx]
        prevent_row_split(row)
        for cidx, value in enumerate(values):
            set_cell_text(row.cells[cidx], value, bold=(ridx == 0), size=9.3)
            set_cell_margins(row.cells[cidx], top=25, bottom=35, start=70, end=110)
            for paragraph in row.cells[cidx].paragraphs:
                paragraph.paragraph_format.keep_with_next = ridx < len(rows) - 1
    set_table_widths(table, [3.35, 3.35])


def add_page_break(doc: Document) -> None:
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def build_sale_agreement() -> Path:
    doc = Document()
    configure_document(doc, "Commercial Machine Sale Agreement")
    add_brand_block(doc, "Commercial Machine Documents")
    add_title(doc, "COMMERCIAL MACHINE SALE AGREEMENT", "Reusable order form, sale terms, and operating exhibits")

    add_paragraph(
        doc,
        "This Commercial Machine Sale Agreement (the “Agreement”) is entered into as of the Effective Date in the Order Form by TGPACI LLC, doing business as Bloomjoy Sweets (“Bloomjoy”), and the buyer identified in the Order Form (“Buyer”). The Order Form, these Standard Terms, and all attached exhibits are one agreement. If they conflict, the Order Form controls for that transaction, followed by the exhibits and then the Standard Terms.",
    )
    p = add_paragraph(doc, "Complete every bracketed field before signature. Delete any option that does not apply.")
    for run in p.runs:
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(PINK)

    add_heading(doc, "PART I — ORDER FORM", level=1)
    add_field_table(
        doc,
        [
            ("Agreement / Quote No.", "[Insert]"),
            ("Effective Date", "[Month Day, Year]"),
            ("Seller", "TGPACI LLC dba Bloomjoy Sweets (“Bloomjoy”)"),
            ("Bloomjoy notice address", "[Insert street address, city, state, ZIP]"),
            ("Bloomjoy notice email", "[Insert legal notice email]"),
            ("Buyer legal name", "[Insert exact legal name and entity type]"),
            ("Buyer notice address", "[Insert street address, city, state, ZIP]"),
            ("Buyer notice email", "[Insert]"),
            ("Buyer billing contact", "[Name | email | phone]"),
            ("Buyer delivery contact", "[Name | email | phone]"),
        ],
    )

    add_heading(doc, "1. Equipment and Price", level=2)
    add_grid_table(
        doc,
        ["Item", "Description / configuration", "Qty.", "Unit price", "Extended price"],
        [
            ["Commercial Machine", "[Model, finish, pattern package, payment reader, connectivity]", "[ ]", "$[ ]", "$[ ]"],
            ["Custom wrap", "☐ Not included   ☐ Included; artwork requirements attached", "[ ]", "$[ ]", "$[ ]"],
            ["Opening supplies", "[Sugar, sticks, tools, spare parts, other]", "[ ]", "$[ ]", "$[ ]"],
            ["Freight / delivery", "[Method, accessorial charges, insurance]", "—", "—", "$[ ]"],
            ["Sales or use tax", "[Estimated; final tax may adjust invoice]", "—", "—", "$[ ]"],
            ["TOTAL PURCHASE PRICE", "", "", "", "$[ ]"],
        ],
        [1.15, 3.05, 0.5, 0.85, 1.1],
    )

    add_heading(doc, "2. Payment Schedule", level=2)
    add_grid_table(
        doc,
        ["Payment", "Amount", "Due"],
        [
            ["Order deposit", "50% of Total Purchase Price: $[ ]", "At signing / purchase"],
            ["Delivery balance", "Remaining 50%, plus approved adjustments: $[ ]", "Within [3] business days after Confirmed Delivery"],
        ],
        [1.35, 2.35, 3.0],
        font_size=9.1,
    )
    add_paragraph(doc, "Payment method / invoice instructions: [Insert].")
    add_paragraph(doc, "Deposit treatment: [Select one] ☐ Refundable until Bloomjoy commits the order to production, less documented nonrecoverable costs. ☐ Nonrefundable upon receipt. ☐ Other: [Insert].")

    add_heading(doc, "3. Delivery and Commissioning", level=2)
    add_field_table(
        doc,
        [
            ("Delivery point", "[Full address, suite/loading instructions]"),
            ("Requested delivery window", "[Insert; estimate only unless expressly guaranteed]"),
            ("Delivery method / Incoterm", "[Carrier / white-glove / customer pickup / other]"),
            ("Included installation support", "[Remote setup; on-site services, if any]"),
            ("Included training", "[Participants, format, duration]"),
            ("Acceptance review period", "[5] business days after Confirmed Delivery"),
            ("Balance timing override", "[Leave blank unless different from Section 2]"),
        ],
    )

    add_heading(doc, "4. Warranty, Support, and Optional Services", level=2)
    add_field_table(
        doc,
        [
            ("Limited warranty term", "[Insert manufacturer-backed term and start date rule]"),
            ("Manufacturer support", "24/7 first-line remote technical support via WeChat, subject to manufacturer availability and policies"),
            ("Bloomjoy included support", "Onboarding and reasonable escalation coordination during U.S. business hours, as described in Exhibit B"),
            ("Bloomjoy Plus", "Not included. Optional, separately purchased, and governed by the online terms described below."),
            ("Additional paid services", "[On-site labor, travel, extended coverage, storage, other—or “None”]"),
        ],
    )

    add_heading(doc, "5. Transaction-Specific Terms", level=2)
    add_field_table(
        doc,
        [
            ("Governing law", "[California, unless another state is inserted]"),
            ("Exclusive venue", "[County and state; default is Santa Clara County, California]"),
            ("Buyer insurance requirement", "[Insert limits, if any]"),
            ("Special terms", "[Insert or attach; write “None” if none]"),
        ],
    )

    add_heading(doc, "6. Signatures", level=2)
    add_paragraph(doc, "Each signer represents that the signer is authorized to bind the identified party. Electronic signatures and counterparts are effective.")
    add_signature_table(doc, "TGPACI LLC dba Bloomjoy Sweets", "BUYER: [Legal name]")

    add_page_break(doc)
    add_heading(doc, "PART II — STANDARD TERMS", level=1)
    terms = [
        (
            "1. Sale and Purchase.",
            "Bloomjoy will sell, and Buyer will purchase, the equipment and included items listed in the Order Form (collectively, the “Equipment”). No product, service, location, territory, permit right, revenue opportunity, or exclusivity is included unless expressly listed. Buyer is acquiring the Equipment for commercial use and has independently evaluated its business plan, site, and expected economics.",
        ),
        (
            "2. Price, Taxes, and Payment.",
            "Buyer will pay the Total Purchase Price on the schedule in the Order Form. The initial payment is a deposit applied to the purchase price; its refundability is governed by the selected deposit treatment. Buyer is responsible for sales, use, excise, and similar transaction taxes, excluding taxes on Bloomjoy’s net income. Approved change orders, carrier adjustments, storage, redelivery, and buyer-caused accessorial charges may be added to the final invoice. Overdue amounts accrue interest at the lesser of 1.0% per month or the maximum lawful rate, plus reasonable collection costs. Buyer may not set off unrelated claims against an amount due.",
        ),
        (
            "3. Order Changes and Cancellation.",
            "Changes require written approval and may affect price and schedule. If Buyer cancels after Bloomjoy or its supplier commits the order, Buyer is responsible for documented nonrecoverable production, customization, freight, storage, and cancellation costs, subject to the deposit treatment in the Order Form. Bloomjoy will return any remaining refundable balance after those amounts are determined.",
        ),
        (
            "4. Delivery; Confirmed Delivery.",
            "Delivery dates are good-faith estimates unless the Order Form expressly states a guaranteed date. “Confirmed Delivery” occurs when the Equipment is physically tendered at the Delivery Point and Buyer signs the Delivery and Commissioning Certificate, or when Buyer receives the Equipment and does not provide a written notice of material shipping damage or a material quantity discrepancy within two business days. A commissioning item that does not prevent delivery does not postpone Confirmed Delivery unless the Order Form expressly says otherwise. Buyer will provide safe access, an authorized recipient, and any required unloading assistance. Re-delivery or storage caused by Buyer is chargeable to Buyer.",
        ),
        (
            "5. Risk of Loss; Title; Security Interest.",
            "Risk of loss passes to Buyer at Confirmed Delivery, except to the extent a carrier claim is controlled by Bloomjoy under the selected delivery method. Title passes only after Bloomjoy receives all amounts due for the Equipment. Until then, Buyer grants Bloomjoy a purchase-money security interest in the Equipment and its proceeds, will keep the Equipment identifiable and free of other liens, and authorizes Bloomjoy to file financing statements reasonably necessary to perfect that interest. Buyer may operate the Equipment in the ordinary course but may not sell, pledge, or relocate it outside the United States without Bloomjoy’s written consent before full payment.",
        ),
        (
            "6. Inspection and Acceptance.",
            "Buyer will inspect promptly. Visible damage, missing items, and material delivery discrepancies must be listed on the Delivery and Commissioning Certificate. Buyer has the Acceptance Review Period to report a material nonconformity in reasonable detail with photographs, video, logs, or other available evidence. Bloomjoy may inspect and, at its option, repair, replace, complete, or credit the affected item. Acceptance does not waive latent defects or an express warranty claim. A minor punch-list item does not permit rejection of the Equipment or withholding of the entire delivery balance.",
        ),
        (
            "7. Site Readiness and Installation.",
            "Buyer is responsible for the site and for completing Exhibit C before delivery, including approved electrical service, stable network connectivity, level floor space, ventilation and clearances, venue authorization, sanitation access, loading access, and all permits not expressly covered by a separate signed agreement. Bloomjoy is not responsible for delay, damage, or service failure caused by an unready site, unauthorized installation, incompatible service, or Buyer’s contractors. Services beyond those listed in the Order Form require a written change order.",
        ),
        (
            "8. Embedded Software and Connectivity.",
            "The Equipment may contain firmware, operating software, design files, cloud connectivity, payment integrations, and third-party services (“Embedded Software”). Subject to full payment and compliance with this Agreement, Buyer receives a nonexclusive, nontransferable right to use the Embedded Software solely with the Equipment for its intended commercial operation. The Embedded Software is licensed, not sold. Buyer may not copy, reverse engineer, bypass controls, introduce malicious code, use unauthorized software that creates a safety or compatibility risk, or use the Embedded Software to develop a competing product, except to the limited extent a restriction is prohibited by law. Manufacturer, payment processor, connectivity, and other third-party terms may also apply. This Section does not require Buyer to purchase Bloomjoy Plus.",
        ),
        (
            "9. Bloomjoy Plus Is Separate and Optional.",
            "Bloomjoy Plus is not a license fee, royalty, or condition of Equipment ownership. Buyer may subscribe or cancel through Bloomjoy’s website at then-current pricing and terms. A cancellation is effective at the end of the paid billing period, and fees are not prorated or refunded unless required by law or the online terms state otherwise. Current information is available at https://www.bloomjoyusa.com/plus, https://www.bloomjoyusa.com/terms, and https://www.bloomjoyusa.com/billing-cancellation. A later Plus subscription does not amend this Agreement.",
        ),
        (
            "10. Permits and Legal Compliance.",
            "Except under a separate Permitted Operations and Commissary Agreement and an active site schedule, Buyer is solely responsible for permits, food-safety approvals, commissary arrangements, taxes, licenses, accessibility, zoning, venue rules, and other legal requirements for possession and operation. No permit is sold, transferred, leased, or granted under this Agreement. Buyer will not represent that Bloomjoy’s permit or commissary agreement covers a site or operator without a separate written activation signed by Bloomjoy and any approvals required by the applicable authority.",
        ),
        (
            "11. Branding and Custom Artwork.",
            "Each party retains its names, logos, and other intellectual property. Buyer grants Bloomjoy and its vendors a limited right to reproduce Buyer-approved artwork solely to produce the ordered wrap or materials. Buyer represents that it has rights to supplied artwork and will indemnify Bloomjoy against a third-party claim arising from that artwork. Buyer may display Bloomjoy branding already affixed to the Equipment but receives no franchise, trademark system, exclusive territory, or right to hold itself out as Bloomjoy’s agent. Any broader brand license must be in a separate writing.",
        ),
        (
            "12. Care, Maintenance, and Records.",
            "Buyer will operate, clean, maintain, and secure the Equipment according to manuals, training, safety notices, and applicable law; use compatible supplies; maintain transaction and service records reasonably needed for support; and promptly stop operation if a condition may threaten health, safety, property, or the Equipment. Buyer is responsible for unauthorized modifications, misuse, vandalism, pests, unsuitable power, network interruption, and ordinary wear unless an express warranty provides otherwise.",
        ),
        (
            "13. Limited Warranty and Support.",
            "The express warranty and support commitments are limited to the Order Form and Exhibit B. Except for those express commitments and to the maximum extent permitted by law, the Equipment, Embedded Software, and services are provided without other warranties, including implied warranties of merchantability, fitness for a particular purpose, noninfringement, uptime, profitability, or suitability of a site. Some jurisdictions do not allow particular exclusions, so an exclusion applies only to the extent lawful.",
        ),
        (
            "14. Confidentiality.",
            "Each party will protect the other’s nonpublic business, technical, pricing, customer, and security information using reasonable care and use it only for this transaction. This duty does not cover information that is public without breach, already lawfully known, independently developed, or rightfully received without restriction. A required disclosure is permitted if the receiving party gives prompt notice when lawful and reasonably limits the disclosure. Trade-secret obligations survive while the information remains a trade secret; other confidentiality obligations survive three years.",
        ),
        (
            "15. Mutual Indemnity.",
            "Each party will defend and indemnify the other and its personnel from a third-party claim to the extent caused by the indemnifying party’s negligence, willful misconduct, violation of law, or material breach. Buyer’s obligation also covers Buyer’s operation, products, venue, employees, contractors, and customer claims, except to the extent caused by a covered defect for which Bloomjoy is responsible. Bloomjoy’s obligation also covers a claim that Bloomjoy-provided branding or materials, used as authorized, infringe a U.S. intellectual-property right. The indemnified party must provide prompt notice, reasonable cooperation, and control of the defense, but no settlement may admit fault or impose a nonmonetary obligation on it without consent.",
        ),
        (
            "16. Limitation of Liability.",
            "To the maximum extent permitted by law, neither party is liable for lost profits, lost revenue, loss of data, business interruption, or indirect, incidental, special, exemplary, or consequential damages arising from this Agreement, even if advised they were possible. Bloomjoy’s aggregate liability arising from the sale will not exceed the Total Purchase Price actually paid. These limits do not apply to payment obligations, misuse of the other party’s intellectual property, breach of confidentiality, indemnification obligations, fraud, gross negligence, willful misconduct, or liability that cannot lawfully be limited.",
        ),
        (
            "17. Default and Remedies.",
            "A party is in material default if it fails to cure a material nonpayment breach within 10 days after written notice or another material breach within 30 days after written notice, or becomes subject to an insolvency proceeding not dismissed within 60 days. If Buyer does not pay the delivery balance when due, Bloomjoy may suspend nonessential services, pursue the Equipment and proceeds as a secured party, or exercise other lawful remedies. Remedies are cumulative. Neither party is required to continue performance when doing so would violate law or create an immediate safety risk.",
        ),
        (
            "18. Force Majeure.",
            "Neither party is liable for delay caused by an event beyond its reasonable control, including supplier interruption, carrier delay, port disruption, labor action, natural disaster, epidemic, government action, utility failure, or cyberattack, if it promptly gives notice and uses commercially reasonable efforts to mitigate. This Section does not excuse payment for Equipment already delivered.",
        ),
        (
            "19. Governing Law and Forum.",
            "The law identified in the Order Form governs without regard to conflict-of-law rules. The parties consent to the exclusive state and federal courts in the venue identified in the Order Form. If either field is blank, California law and the courts located in Santa Clara County, California apply. Before filing a claim other than for emergency injunctive relief or collection of an undisputed amount, an executive from each party will attempt in good faith to resolve the dispute for at least 15 days after written notice.",
        ),
        (
            "20. General.",
            "The parties are independent contractors. Neither may bind the other. Buyer may not assign this Agreement before full payment without Bloomjoy’s written consent; Bloomjoy may assign it to an affiliate or in connection with a sale of substantially all relevant assets. Notices must be sent to the Order Form contacts by personal delivery, nationally recognized overnight service, certified mail, or email with confirmation of receipt. Waivers must be written. If a provision is unenforceable, it will be modified to the minimum extent necessary and the remainder will continue. This Agreement is the complete agreement for the sale and may be amended only in a signed writing, except that applicable third-party service terms may update under their own terms. Headings are for convenience. Counterparts and electronic signatures are valid. Sections that by their nature should survive will survive, including payment, title and security interest, intellectual property, confidentiality, indemnity, liability limits, and dispute terms.",
        ),
    ]
    # Keep the dense standard terms readable and balanced across three pages.
    add_numbered_terms(doc, terms[:7])
    add_page_break(doc)
    add_numbered_terms(doc, terms[7:14])
    add_page_break(doc)
    add_numbered_terms(doc, terms[14:])

    add_page_break(doc)
    add_heading(doc, "EXHIBIT A — EQUIPMENT AND DELIVERY SPECIFICATIONS", level=1)
    add_field_table(
        doc,
        [
            ("Agreement / Quote No.", "[Insert]"),
            ("Machine model", "[Insert]"),
            ("Quantity", "[Insert]"),
            ("Serial number(s)", "[Assign before or at delivery]"),
            ("Cabinet finish / wrap", "[Insert]"),
            ("Payment reader / processor", "[Insert; merchant account owner]"),
            ("Connectivity", "[Wi-Fi / cellular / Ethernet; owner of service]"),
            ("Electrical requirements", "[Voltage, amperage, outlet, dedicated circuit]"),
            ("Dimensions / weight", "[Insert confirmed specifications]"),
            ("Included accessories", "[Insert]"),
            ("Included opening supplies", "[Insert]"),
            ("Documentation / credentials", "[Manuals, admin access, support channel]"),
            ("Packaging / freight", "[Insert]"),
            ("Unloading responsibility", "[Bloomjoy / Buyer / carrier]"),
            ("Installation responsibility", "[Insert]"),
            ("Commissioning test", "[Insert required test and responsible person]"),
        ],
    )
    add_heading(doc, "Approved Attachments", level=2)
    add_bullets(
        doc,
        [
            "☐ Final quote or invoice",
            "☐ Final wrap proof and approved artwork",
            "☐ Confirmed manufacturer specification sheet",
            "☐ Freight quote and delivery instructions",
            "☐ Additional written scope or change order",
        ],
    )

    add_page_break(doc)
    add_heading(doc, "EXHIBIT B — LIMITED WARRANTY AND SUPPORT", level=1)
    add_numbered_terms(
        doc,
        [
            ("1. Warranty Term.", "The limited warranty begins on [Confirmed Delivery / successful commissioning] and continues for the period stated in the Order Form. Any “up to” public warranty statement is not a substitute for the completed Order Form."),
            ("2. Covered Defects.", "During the warranty term, Bloomjoy will coordinate the manufacturer-backed repair or replacement process for a material defect in parts or workmanship under normal intended use. The remedy may include remote diagnosis, software adjustment, shipment of a replacement part, or another commercially reasonable correction. Replaced parts may be new or functionally equivalent."),
            ("3. Exclusions.", "The warranty does not cover consumables, ordinary wear, cosmetic conditions that do not impair operation, improper cleaning, unauthorized repair or modification, incompatible supplies, vandalism, pests, accident, unsuitable power or network service, environmental conditions, relocation damage, failure to follow manuals or safety instructions, or a third-party payment or connectivity service."),
            ("4. Claim Process.", "Buyer will stop use if continued operation may worsen damage or create a safety risk; contact the designated manufacturer support channel; provide the serial number, photographs or video, logs, and requested diagnostic information; and reasonably cooperate with remote troubleshooting. Buyer must notify Bloomjoy at [support email] if escalation or parts coordination is needed."),
            ("5. Labor, Travel, and Shipping.", "Included labor, freight, duties, travel, and on-site services are limited to: [Insert]. Any amount not expressly included requires Buyer approval before charge, except reasonable emergency measures requested by Buyer."),
            ("6. Support Boundaries.", "The manufacturer provides first-line technical support through its designated 24/7 WeChat channel, subject to actual availability, time zone, and issue context. Bloomjoy provides onboarding guidance and reasonable escalation coordination during U.S. business hours. Bloomjoy does not promise continuous uptime, a fixed response time, or on-site service unless the Order Form states otherwise."),
            ("7. Bloomjoy Plus.", "Optional Bloomjoy Plus training, playbooks, portal features, or concierge benefits are governed only by the online subscription terms. Subscription status does not expand or reduce the express machine warranty unless a signed order expressly says so."),
        ],
    )

    add_page_break(doc)
    add_heading(doc, "EXHIBIT C — SITE READINESS CHECKLIST", level=1)
    add_paragraph(doc, "Buyer must complete and return this checklist no later than [10] business days before the requested delivery date. A checked item confirms completion, not merely an intention to complete it.")
    add_grid_table(
        doc,
        ["Ready", "Requirement", "Owner / notes"],
        [
            ["☐", "Final placement and venue authorization confirmed", "[ ]"],
            ["☐", "Required permits, licenses, and food-safety approvals confirmed", "[ ]"],
            ["☐", "Level floor space, clearances, guest flow, and service access confirmed", "[ ]"],
            ["☐", "Correct dedicated electrical service and surge protection available", "[ ]"],
            ["☐", "Stable network / cellular service tested at placement", "[ ]"],
            ["☐", "Payment account, reader, and transaction settlement plan ready", "[ ]"],
            ["☐", "Loading path, dock/elevator, doorway dimensions, and unloading labor confirmed", "[ ]"],
            ["☐", "Cleaning, handwashing, water, waste, pest-control, and sanitation plan ready", "[ ]"],
            ["☐", "Approved sugar, sticks, tools, and initial stock available", "[ ]"],
            ["☐", "Authorized operators identified and training scheduled", "[ ]"],
            ["☐", "Insurance and site-required certificates delivered", "[ ]"],
            ["☐", "Secure storage and after-hours access plan confirmed", "[ ]"],
        ],
        [0.55, 4.4, 1.75],
        font_size=8.8,
    )
    add_field_table(
        doc,
        [
            ("Known exceptions", "[Insert or “None”]"),
            ("Buyer readiness contact", "[Name | email | phone]"),
            ("Buyer certification", "I certify that the information above is accurate and will promptly report any change before delivery."),
            ("Name / title", "[Insert]"),
            ("Signature / date", "__________________________________   __________________"),
        ],
    )

    path = OUT_DIR / "Commercial Machine Sale Agreement.docx"
    doc.save(path)
    return path


def build_delivery_certificate() -> Path:
    doc = Document()
    configure_document(doc, "Delivery and Commissioning Certificate")
    add_brand_block(doc, "Commercial Machine Documents")
    add_title(doc, "DELIVERY AND COMMISSIONING CERTIFICATE", "Use for each machine delivered under a Commercial Machine Sale Agreement")
    add_paragraph(doc, "This certificate documents delivery condition, commissioning results, open items, and the date that triggers the delivery balance under the referenced sale agreement. It does not replace the sale agreement or waive a latent defect or express warranty claim.")

    add_heading(doc, "1. Transaction and Delivery Details", level=1)
    add_field_table(
        doc,
        [
            ("Sale Agreement / Quote No.", "[Insert]"),
            ("Buyer legal name", "[Insert]"),
            ("Delivery point", "[Insert]"),
            ("Delivery date and time", "[Insert]"),
            ("Machine model", "[Insert]"),
            ("Serial number", "[Insert]"),
            ("Carrier / delivery provider", "[Insert]"),
            ("Person receiving delivery", "[Name | title | phone]"),
        ],
    )

    add_heading(doc, "2. Delivery Condition", level=1)
    add_grid_table(
        doc,
        ["Inspection item", "Pass", "Follow-up", "N/A", "Notes / evidence"],
        [
            ["Packaging and shock / tilt indicators", "☐", "☐", "☐", "[ ]"],
            ["Cabinet, glass, doors, locks, and exterior", "☐", "☐", "☐", "[ ]"],
            ["Machine model and serial match order", "☐", "☐", "☐", "[ ]"],
            ["Accessories, tools, manuals, and credentials", "☐", "☐", "☐", "[ ]"],
            ["Opening supplies and spare parts", "☐", "☐", "☐", "[ ]"],
            ["Custom wrap / branding matches approved proof", "☐", "☐", "☐", "[ ]"],
            ["Visible freight damage photographed and noted", "☐", "☐", "☐", "[ ]"],
        ],
        [2.55, 0.55, 0.78, 0.5, 2.32],
        font_size=8.4,
    )

    add_heading(doc, "3. Commissioning Checks", level=1)
    add_grid_table(
        doc,
        ["Commissioning item", "Pass", "Follow-up", "N/A", "Notes / evidence"],
        [
            ["Correct power and safe startup", "☐", "☐", "☐", "[ ]"],
            ["Display, controls, emergency stop, and doors", "☐", "☐", "☐", "[ ]"],
            ["Stick dispenser / handoff operates", "☐", "☐", "☐", "[ ]"],
            ["Sugar bins / feed path configured", "☐", "☐", "☐", "[ ]"],
            ["Burner / spinning head / production cycle tested", "☐", "☐", "☐", "[ ]"],
            ["At least one complete product cycle observed", "☐", "☐", "☐", "[ ]"],
            ["Network connection and time zone configured", "☐", "☐", "☐", "[ ]"],
            ["Payment reader test transaction completed", "☐", "☐", "☐", "[ ]"],
            ["Dashboard / remote visibility confirmed, if included", "☐", "☐", "☐", "[ ]"],
            ["Manuals and manufacturer support channel provided", "☐", "☐", "☐", "[ ]"],
            ["Operator received startup / shutdown orientation", "☐", "☐", "☐", "[ ]"],
        ],
        [2.55, 0.55, 0.78, 0.5, 2.32],
        font_size=8.4,
    )

    add_heading(doc, "4. Exceptions and Open Items", level=1)
    add_grid_table(
        doc,
        ["No.", "Issue / evidence", "Owner", "Target date", "Resolution / close date"],
        [
            ["1", "[Insert or write “None”]", "[ ]", "[ ]", "[ ]"],
            ["2", "[ ]", "[ ]", "[ ]", "[ ]"],
            ["3", "[ ]", "[ ]", "[ ]", "[ ]"],
            ["4", "[ ]", "[ ]", "[ ]", "[ ]"],
        ],
        [0.45, 2.65, 1.05, 1.1, 1.45],
        font_size=8.7,
    )

    add_heading(doc, "5. Delivery Status", level=1)
    add_bullets(
        doc,
        [
            "☐ Accepted — Confirmed Delivery occurred on the date above.",
            "☐ Accepted with listed open items — Confirmed Delivery occurred; listed items will be handled under the sale agreement and do not prevent ordinary use except as noted.",
            "☐ Material delivery nonconformity — delivery is not accepted for the specific reasons listed below.",
        ],
    )
    add_field_table(
        doc,
        [
            ("Material nonconformity details", "[Insert only if the third box is selected]"),
            ("Safe to operate pending follow-up?", "☐ Yes   ☐ No   ☐ With restrictions: [Insert]"),
            ("Confirmed Delivery date", "[Insert]"),
            ("Warranty start date", "[Insert according to the sale agreement]"),
        ],
    )

    add_heading(doc, "6. Payment Acknowledgment", level=1)
    add_paragraph(doc, "If delivery is accepted or accepted with listed open items, Buyer acknowledges that Confirmed Delivery has occurred and the remaining 50% delivery balance is due on the schedule stated in the sale agreement. This acknowledgment does not waive a timely claim for concealed freight damage, a latent defect, or an express warranty remedy.")

    add_heading(doc, "7. Signatures", level=1)
    add_signature_table(doc, "BUYER’S AUTHORIZED REPRESENTATIVE", "BLOOMJOY REPRESENTATIVE")
    add_paragraph(doc, "Carrier / delivery provider acknowledgment (optional):")
    add_field_table(
        doc,
        [
            ("Name / company", "[Insert]"),
            ("Signature / date", "__________________________________   __________________"),
        ],
    )

    path = OUT_DIR / "Delivery and Commissioning Certificate.docx"
    doc.save(path)
    return path


def build_permitted_operations_agreement() -> Path:
    doc = Document()
    configure_document(doc, "Permitted Operations and Commissary Agreement")
    add_brand_block(doc, "Optional Operations Program")
    add_title(doc, "PERMITTED OPERATIONS AND COMMISSARY AGREEMENT", "Optional 6% program; activate each location with a separate Site Activation Schedule")

    add_paragraph(
        doc,
        "This Permitted Operations and Commissary Agreement (the “Program Agreement”) is entered into as of [Month Day, Year] by TGPACI LLC, doing business as Bloomjoy Sweets (“Bloomjoy”), and [Operator legal name] (“Operator”). It is separate from any equipment sale and from Bloomjoy Plus. It does not sell, lease, transfer, or sublicense a health permit. Each site may operate under this Program Agreement only when the applicable authority permits the proposed structure and the parties sign a complete Site Activation Schedule.",
    )
    p = add_paragraph(doc, "Do not activate a site until every activation condition in Section 3 and Exhibit A is complete.")
    for run in p.runs:
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(PINK)

    add_heading(doc, "PARTIES AND PROGRAM INFORMATION", level=1)
    add_field_table(
        doc,
        [
            ("Effective Date", "[Insert]"),
            ("Bloomjoy notice address", "[Insert street address, city, state, ZIP]"),
            ("Bloomjoy notice email", "[Insert legal notice email]"),
            ("Operator legal name", "[Insert exact legal name and entity type]"),
            ("Operator notice address", "[Insert street address, city, state, ZIP]"),
            ("Operator notice email", "[Insert]"),
            ("Program contact", "[Name | email | phone]"),
            ("Default Program Fee", "6% of Net Sales for each Activated Site, unless its Site Activation Schedule states otherwise"),
            ("Initial term", "[12 months / month-to-month / other]"),
        ],
    )

    add_heading(doc, "PROGRAM TERMS", level=1)
    terms = [
        (
            "1. Purpose and Structure.",
            "The Program Agreement establishes a controlled, site-specific arrangement under which Operator may perform approved machine-operating services and may use a listed commissary or other approved facility, only to the extent authorized in an Active Site Schedule and by the applicable enforcement agency. The legal operator, permit holder, merchant of record, products, location, machine, and commissary are identified separately for each site. A Site Activation Schedule is “Active” only after all required signatures and written approvals are complete.",
        ),
        (
            "2. No Permit Transfer or Blanket Coverage.",
            "A food-facility or health permit remains with the person, location, activity, and time period for which the authority issued it. Nothing in this Program Agreement transfers, rents, sublicenses, or extends a permit. Operator will not use Bloomjoy’s name, permit number, commissary agreement, or regulatory relationship for a location or activity not expressly listed in an Active Site Schedule. A sale of Equipment and payment of the Program Fee do not create permit coverage.",
        ),
        (
            "3. Conditions to Site Activation.",
            "Before operation begins, the parties must complete Exhibit A; identify the permit holder and legal operating structure; obtain written confirmation or approval from the local enforcement agency when Bloomjoy reasonably requires it; obtain the commissary’s written acknowledgment in Exhibit B or an equivalent document; confirm venue approval, equipment identification, approved products, food-safety procedures, trained personnel, insurance, payment processing, tax handling, records, and emergency contacts; and satisfy any additional agency condition. Bloomjoy may decline activation if the structure is unclear or exposes either party to unreasonable regulatory, safety, or reputational risk.",
        ),
        (
            "4. Bloomjoy Responsibilities.",
            "For each Active Site, Bloomjoy will perform only the responsibilities assigned to it in the Site Activation Schedule, which may include maintaining a listed permit or commissary agreement, communicating with an enforcement agency or commissary, providing approved operating procedures, reviewing required records, and coordinating inspections. Bloomjoy does not guarantee approval, uninterrupted permit status, or the continued availability of a particular commissary. Bloomjoy will give Operator reasonably prompt notice of a known suspension, restriction, or material change affecting an Active Site.",
        ),
        (
            "5. Operator Responsibilities.",
            "Operator will use only the identified Equipment at the approved location; use only approved ingredients, packaging, and processes; follow Bloomjoy’s and the authority’s food-safety, cleaning, temperature, labeling, allergen, handwashing, waste, pest-control, and incident procedures; ensure only trained and authorized personnel operate or service the Equipment; maintain accurate daily logs and transaction records; allow inspections; promptly correct deficiencies; and immediately report a complaint, illness allegation, contamination concern, pest issue, equipment hazard, agency contact, data-security event, or material venue change. Operator may not relocate the Equipment or materially change the menu, process, owner, or site configuration without written reactivation.",
        ),
        (
            "6. Commissary Access and Use.",
            "Commissary rights are limited to the facility, services, hours, storage, cleaning, waste, and access rules stated in the Active Site Schedule and Exhibit B. Operator receives no tenancy, ownership, or independent contract right unless the commissary separately agrees. Operator will sign in, keep areas sanitary, protect keys and access codes, use only assigned storage, remove waste as directed, and pay the fees allocated to Operator. Bloomjoy may suspend access immediately for a safety, security, payment, or compliance concern.",
        ),
        (
            "7. Program Fee and Net Sales.",
            "Operator will pay Bloomjoy the Program Fee in each Active Site Schedule, with 6% of Net Sales as the default. “Net Sales” means all amounts actually collected from sales through the Equipment or at the Active Site, less only: (a) sales or use taxes separately stated and actually remitted; (b) customer refunds actually paid; and (c) payment chargebacks actually assessed. Net Sales are not reduced by card-processing fees, venue rent or revenue share, labor, insurance, commissary fees, consumables, repairs, delivery costs, or other operating expenses unless the Site Activation Schedule expressly permits the deduction. The fee is compensation for the program services and access stated here; it is not a machine software license fee or a Bloomjoy Plus subscription.",
        ),
        (
            "8. Merchant of Record; Settlement and Taxes.",
            "Each Active Site Schedule identifies the merchant of record and settlement flow. If Operator collects revenue, Operator will deliver the monthly report in Exhibit C and pay the Program Fee within 15 days after month-end. If Bloomjoy collects revenue, Bloomjoy may deduct the Program Fee, processor fees, refunds, chargebacks, taxes, and any other schedule-authorized amounts before remitting the balance on the schedule stated in Exhibit A. The Site Activation Schedule allocates responsibility for sales-tax registration, collection, filing, and remittance. Each party remains responsible for its income and payroll taxes.",
        ),
        (
            "9. Reports, Records, and Audit.",
            "Operator will maintain complete sales, settlement, cleaning, commissary, training, ingredient-lot, maintenance, complaint, and inspection records for at least three years, or longer if law requires. Bloomjoy may access machine or processor data made available for the Active Site and may inspect relevant records on 10 business days’ notice, or immediately for a safety or regulatory issue. If an audit finds an underpayment greater than 5%, Operator will promptly pay the shortfall, lawful interest, and Bloomjoy’s reasonable audit cost. Operator will preserve records during any open claim or agency matter.",
        ),
        (
            "10. Regulatory Direction and Suspension.",
            "Operator will follow a lawful direction from Bloomjoy concerning food safety, permit conditions, commissary use, records, branding required by the authority, or suspension of an Active Site. Bloomjoy may suspend operation immediately if an approval expires or is questioned; a required record is missing; a payment is overdue; an authority, venue, commissary, insurer, or payment provider objects; or Bloomjoy reasonably believes continued operation may create a legal, health, safety, or reputational risk. Suspension does not transfer operating authority to Operator and does not waive accrued fees.",
        ),
        (
            "11. Equipment; Support; Bloomjoy Plus.",
            "Equipment ownership, warranty, and embedded software rights are governed by the applicable sale agreement or other ownership document. This Program Agreement does not modify them. Bloomjoy Plus remains an optional, separately purchased online subscription and is not included in the Program Fee. Termination of Bloomjoy Plus does not itself terminate this Program Agreement, and termination of this Program Agreement does not itself cancel Bloomjoy Plus.",
        ),
        (
            "12. Branding; No Territory or Franchise Representation.",
            "Operator may use only the names and marks specifically approved for an Active Site and only while that schedule remains active. Operator receives no exclusive territory and may not represent itself as the owner of Bloomjoy, a Bloomjoy employee, or the holder of a Bloomjoy permit. The parties intend a limited services and compliance arrangement, not the sale of a franchise or business opportunity. This statement does not override any law; the parties will amend or discontinue a structure if counsel or an authority determines another legal framework is required.",
        ),
        (
            "13. Insurance.",
            "Operator will maintain the insurance listed in each Active Site Schedule, issued by reputable carriers, and provide certificates before activation. Unless the schedule states otherwise, required coverage includes commercial general liability with product/completed-operations coverage, workers’ compensation as required by law, automobile coverage for business transport, and property coverage for Operator-owned Equipment. Bloomjoy and any required venue or commissary party will be named as additional insureds when commercially available and stated in the schedule.",
        ),
        (
            "14. Mutual Indemnity.",
            "Each party will defend and indemnify the other and its personnel from a third-party claim to the extent caused by the indemnifying party’s negligence, willful misconduct, violation of law, or breach. Operator’s obligation includes claims arising from Operator’s personnel, unsafe or unapproved operation, food handling, employment practices, venue conduct, taxes allocated to Operator, or use outside an Active Site. Bloomjoy’s obligation includes claims arising from Bloomjoy’s personnel or Bloomjoy’s failure to perform a responsibility expressly assigned to it. Prompt notice, reasonable cooperation, control of defense, and consent to any settlement imposing fault or nonmonetary duties are required.",
        ),
        (
            "15. Limitation of Liability.",
            "To the maximum extent permitted by law, neither party is liable for lost profits, lost revenue, business interruption, or indirect, incidental, special, exemplary, or consequential damages. Bloomjoy’s aggregate liability under this Program Agreement will not exceed the Program Fees paid or payable for the affected Active Site during the 12 months before the event. These limits do not apply to payment and settlement obligations, misuse of intellectual property or permit information, breach of confidentiality, indemnification obligations, fraud, gross negligence, willful misconduct, or liability that cannot lawfully be limited.",
        ),
        (
            "16. Term; Site Deactivation; Termination.",
            "This Program Agreement begins on the Effective Date and continues for the initial term stated above, then renews month-to-month unless the parties select another term. Either party may terminate the Program Agreement or a Site Activation Schedule on 30 days’ written notice. A party may terminate for an uncured material breach after 10 days for nonpayment or 30 days for another breach. Bloomjoy may immediately deactivate a site under Section 10. On deactivation, Operator will stop using the permit information, commissary access, and program branding; remove or revise required public-facing identifications as directed; return access items and records; pay accrued amounts; and cooperate in a safe wind-down. Deactivation does not prevent Operator from seeking its own approvals and operating independently once lawfully authorized.",
        ),
        (
            "17. Confidentiality and Data.",
            "Each party will reasonably protect the other’s nonpublic business, customer, pricing, permit, security, and technical information and use it only for this Program Agreement. The parties may share information with the authority, commissary, venue, insurer, processor, professional advisor, or service provider as reasonably necessary. Each party will comply with applicable privacy and payment-data rules for data it controls and promptly report a known breach affecting the other’s information.",
        ),
        (
            "18. Independent Contractors; No General Agency.",
            "The parties are independent contractors. Operator has no authority to bind Bloomjoy, incur debt for Bloomjoy, or make a representation on Bloomjoy’s behalf. Any operational direction or limited regulatory role described in an Active Site Schedule exists only to satisfy program and legal requirements and does not create a general employment, partnership, joint venture, or agency relationship. Operator is solely responsible for its employees and contractors, including pay, scheduling, supervision, benefits, and payroll obligations.",
        ),
        (
            "19. Governing Law; Notices; General.",
            "California law governs, without regard to conflict-of-law rules, and the parties consent to the exclusive state and federal courts in Santa Clara County, California, unless an Active Site Schedule states another lawful forum. Notices must be sent to the addresses above by personal delivery, nationally recognized overnight service, certified mail, or email with confirmation of receipt. Neither party may assign this Program Agreement without written consent, except Bloomjoy may assign it to an affiliate or successor to the relevant business. Waivers and amendments must be written. If a provision is unenforceable, it will be narrowed and the remainder will continue. This Program Agreement and its Active Site Schedules are the complete agreement on this program. Electronic signatures and counterparts are valid. Accrued payment, records, confidentiality, indemnity, liability, and dispute terms survive.",
        ),
    ]
    add_numbered_terms(doc, terms)

    add_heading(doc, "SIGNATURES", level=1)
    add_signature_table(doc, "TGPACI LLC dba Bloomjoy Sweets", "OPERATOR: [Legal name]")

    add_page_break(doc)
    add_heading(doc, "EXHIBIT A — SITE ACTIVATION SCHEDULE", level=1)
    add_paragraph(doc, "This Site Activation Schedule is part of the Program Agreement. It becomes Active only when every required field is complete, all activation conditions are satisfied, and both parties sign below.")
    add_field_table(
        doc,
        [
            ("Schedule No. / activation date", "[Insert]"),
            ("Operator", "[Legal name]"),
            ("Site / venue name", "[Insert]"),
            ("Approved operating address", "[Full location, suite, placement description]"),
            ("Local enforcement agency", "[Agency name and contact]"),
            ("Legal permit holder", "[Bloomjoy / Operator / other approved person]"),
            ("Permit number / expiration", "[Insert or “pending”]"),
            ("Approved facility type / activity", "[Insert exactly as approved]"),
            ("Written agency approval", "[Date, contact, attached document]"),
            ("Machine model / serial", "[Insert]"),
            ("Owner of Equipment", "[Insert]"),
            ("Approved products / ingredients", "[Insert]"),
            ("Named responsible person", "[Name | certificate(s) | contact]"),
            ("Authorized operators", "[Names or attached roster]"),
            ("Venue approval / term", "[Document and dates]"),
        ],
    )

    add_heading(doc, "Commissary and Service Plan", level=2)
    add_field_table(
        doc,
        [
            ("Commissary / approved facility", "[Legal name and full address]"),
            ("Commissary permit / contact", "[Number | contact | phone]"),
            ("Authorized services", "☐ Cleaning   ☐ Water   ☐ Waste   ☐ Ingredient storage   ☐ Equipment storage   ☐ Other: [ ]"),
            ("Access days / hours", "[Insert]"),
            ("Storage assignment", "[Insert]"),
            ("Reporting frequency", "[Daily / weekly / other]"),
            ("Commissary fee responsibility", "[Bloomjoy / Operator / included / other]"),
            ("Commissary acknowledgment", "☐ Exhibit B attached   ☐ Equivalent written acknowledgment attached"),
        ],
    )

    add_heading(doc, "Commercial and Settlement Terms", level=2)
    add_field_table(
        doc,
        [
            ("Program Fee", "6% of Net Sales [or insert approved alternative]"),
            ("Merchant of record", "[Bloomjoy / Operator / other]"),
            ("Processor / account owner", "[Insert]"),
            ("Settlement flow", "[Who receives gross proceeds; permitted deductions; remittance recipient]"),
            ("Report due", "15 days after month-end [or insert]"),
            ("Payment / remittance due", "15 days after month-end [or insert]"),
            ("Sales-tax responsibility", "[Registration, collection, filing, remittance]"),
            ("Additional permitted Net Sales deductions", "[Insert or “None”]"),
            ("Site-specific costs", "[Commissary, venue, processor, insurance, permits, other]"),
        ],
    )

    add_heading(doc, "Insurance and Special Conditions", level=2)
    add_field_table(
        doc,
        [
            ("General liability", "$[ ] per occurrence / $[ ] aggregate"),
            ("Product / completed operations", "[Included / separate limit]"),
            ("Workers’ compensation", "[Statutory / exemption evidence]"),
            ("Automobile / property", "[Insert]"),
            ("Additional insureds", "[Bloomjoy, venue, commissary, other]"),
            ("Schedule term / termination", "[Insert if different from Program Agreement]"),
            ("Special agency / venue conditions", "[Insert]"),
        ],
    )

    add_heading(doc, "Activation Checklist", level=2)
    add_bullets(
        doc,
        [
            "☐ Written agency confirmation or approval attached",
            "☐ Commissary acknowledgment attached",
            "☐ Venue authorization attached",
            "☐ Certificate(s) of insurance attached",
            "☐ Responsible person and operator training verified",
            "☐ Cleaning, service, and recordkeeping plan approved",
            "☐ Payment processor and test transaction verified",
            "☐ Public-facing permit holder / operator identification approved",
            "☐ Emergency and incident contacts exchanged",
        ],
    )
    add_signature_table(doc, "BLOOMJOY — SITE ACTIVATION APPROVAL", "OPERATOR — SITE ACCEPTANCE")

    add_page_break(doc)
    add_heading(doc, "EXHIBIT B — COMMISSARY ACKNOWLEDGMENT", level=1)
    add_paragraph(doc, "The commissary or approved facility completes this acknowledgment for the site identified below. This form does not replace a document required by the local enforcement agency.")
    add_field_table(
        doc,
        [
            ("Commissary legal / business name", "[Insert]"),
            ("Facility address", "[Insert]"),
            ("Permit number / agency", "[Insert]"),
            ("Commissary representative", "[Name | title | email | phone]"),
            ("Bloomjoy permit / program contact", "[Insert]"),
            ("Operator", "[Insert]"),
            ("Active Site / machine", "[Insert site and serial]"),
            ("Authorized services", "☐ Cleaning   ☐ Water   ☐ Waste   ☐ Ingredient storage   ☐ Equipment storage   ☐ Other: [ ]"),
            ("Access schedule", "[Insert]"),
            ("Assigned storage / restrictions", "[Insert]"),
            ("Fees / payer", "[Insert]"),
            ("Start / end date", "[Insert]"),
            ("Required logs", "[Insert]"),
            ("Special conditions", "[Insert or “None”]"),
        ],
    )
    add_paragraph(doc, "The undersigned acknowledges the listed access and services, subject to facility rules, permit conditions, and the right to suspend access for a safety, regulatory, security, or payment issue. The undersigned will notify Bloomjoy and Operator of a material change that affects the listed arrangement.")
    add_signature_table(doc, "COMMISSARY / APPROVED FACILITY", "BLOOMJOY ACKNOWLEDGMENT")
    add_paragraph(doc, "Operator acknowledgment:")
    add_field_table(
        doc,
        [
            ("Name / title", "[Insert]"),
            ("Signature / date", "__________________________________   __________________"),
        ],
    )

    add_page_break(doc)
    add_heading(doc, "EXHIBIT C — MONTHLY SALES AND COMPLIANCE REPORT", level=1)
    add_field_table(
        doc,
        [
            ("Reporting month", "[Month / Year]"),
            ("Site / Schedule No.", "[Insert]"),
            ("Machine serial", "[Insert]"),
            ("Merchant of record", "[Insert]"),
            ("Prepared by / date", "[Name | title | date]"),
        ],
    )
    add_heading(doc, "Sales Reconciliation", level=2)
    add_grid_table(
        doc,
        ["Line", "Description", "Amount"],
        [
            ["A", "Gross amounts actually collected", "$[ ]"],
            ["B", "Less sales / use taxes separately stated and actually remitted", "($[ ])"],
            ["C", "Less customer refunds actually paid", "($[ ])"],
            ["D", "Less payment chargebacks actually assessed", "($[ ])"],
            ["E", "Less other deduction expressly allowed by Site Activation Schedule", "($[ ])"],
            ["F", "NET SALES (A − B − C − D − E)", "$[ ]"],
            ["G", "Program Fee percentage", "6% [or schedule rate]"],
            ["H", "PROGRAM FEE DUE (F × G)", "$[ ]"],
            ["I", "Other approved settlement adjustment", "$[ ]"],
            ["J", "Amount paid / remitted with report", "$[ ]"],
        ],
        [0.55, 4.75, 1.4],
        font_size=8.9,
    )
    add_heading(doc, "Compliance Certification", level=2)
    add_grid_table(
        doc,
        ["Confirm", "Monthly requirement", "Notes / attachment"],
        [
            ["☐", "Daily cleaning and service logs complete", "[ ]"],
            ["☐", "Commissary visits / required logs complete", "[ ]"],
            ["☐", "Only approved products, ingredients, and operators used", "[ ]"],
            ["☐", "No unreported relocation or material operating change", "[ ]"],
            ["☐", "No customer illness, contamination, pest, injury, or agency incident—or all incidents attached", "[ ]"],
            ["☐", "Processor statement and tax support attached or available", "[ ]"],
            ["☐", "Open corrective actions identified below", "[ ]"],
        ],
        [0.62, 4.55, 1.53],
        font_size=8.7,
    )
    add_field_table(
        doc,
        [
            ("Incidents / corrective actions", "[Insert or “None”]"),
            ("Certification", "I certify that this report is complete and accurate and that supporting records will be retained as required by the Program Agreement."),
            ("Operator name / title", "[Insert]"),
            ("Signature / date", "__________________________________   __________________"),
            ("Bloomjoy review / date", "[Insert]"),
        ],
    )

    path = OUT_DIR / "Permitted Operations and Commissary Agreement.docx"
    doc.save(path)
    return path


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    paths = [
        build_sale_agreement(),
        build_delivery_certificate(),
        build_permitted_operations_agreement(),
    ]
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
